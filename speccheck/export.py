"""Write SpecCheck results: versioned Excel test library, traceability matrix, Jira CSV, review report and
Playwright test skeletons."""

import csv
import io
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from speccheck.generate import Drafted, TestCase
from speccheck.lint import Gap
from speccheck.parse import Story

HEAD = PatternFill("solid", fgColor="D9E1F2")
MISSING = PatternFill("solid", fgColor="FCE4E4")
ASSUMED = PatternFill("solid", fgColor="FFF4E0")


def case_id(story: Story, n: int) -> str:
    return f"TC-{story.id.split('-', 1)[1]}-{n:02d}"


@dataclass
class Trace:
    story: str
    number: int
    criterion: str
    cases: list[str]
    gaps: list[str]

    @property
    def coverage(self) -> str:
        """NO: no test. assumed: tested, but the criterion has an open gap, so the expected result was assumed."""
        return "NO" if not self.cases else "assumed" if self.gaps else "yes"


def traceability(stories: list[Story], drafts: dict[str, Drafted], gaps: list[Gap]) -> list[Trace]:
    """Acceptance criterion → test cases that claim to cover it, plus open gaps on that criterion."""
    rows = []
    for s in stories:
        cases = drafts.get(s.id, Drafted(test_cases=[])).test_cases
        for i, c in enumerate(s.criteria, start=1):
            rows.append(Trace(s.id, i, c, [case_id(s, n) for n, tc in enumerate(cases, start=1) if i in tc.covers],
                              [g.rule for g in gaps if g.story == s.id and g.text == c]))
    return rows


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), HEAD
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def excel_library(path, version: str, prd_title: str, stories: list[Story], drafts: dict[str, Drafted],
                  gaps: list[Gap], previous: Path | None = None):
    """Write the library to `path` (a file path or a binary stream) and return it."""
    wb = Workbook()
    wb.remove(wb.active)
    cases = []
    for s in stories:
        for n, tc in enumerate(drafts.get(s.id, Drafted(test_cases=[])).test_cases, start=1):
            cases.append([case_id(s, n), s.id, ", ".join(map(str, tc.covers)), tc.title, tc.type, tc.priority,
                          tc.preconditions, "\n".join(f"{i}. {st}" for i, st in enumerate(tc.steps, start=1)),
                          tc.expected, version, "draft", "", ""])
    _sheet(wb, "Test cases", ["ID", "Story", "Covers AC", "Title", "Type", "Priority", "Preconditions", "Steps",
                              "Expected result", "Added in", "Status", "Executed by / date", "Actual result"],
           cases, [11, 8, 8, 40, 11, 9, 30, 55, 45, 9, 9, 18, 30])
    trace = traceability(stories, drafts, gaps)
    ws = _sheet(wb, "Traceability", ["Story", "AC #", "Acceptance criterion", "Test cases", "Open gaps", "Covered"],
                [[t.story, t.number, t.criterion, ", ".join(t.cases), ", ".join(t.gaps), t.coverage]
                 for t in trace], [8, 6, 60, 30, 30, 9])
    for row in ws.iter_rows(min_row=2):
        if row[5].value in ("NO", "assumed"):
            for cell in row:
                cell.fill = MISSING if row[5].value == "NO" else ASSUMED
    gap_rows = [[f"GAP-{i:02d}", g.story, g.severity, g.rule, g.text, g.why, g.question, "open"]
                for i, g in enumerate(gaps, start=1)]
    gap_rows += [[f"GAP-L{i:02d}", sid, "medium", "llm_ambiguity", a, "Drafting could not state an exact expected result.",
                  "Please clarify.", "open"]
                 for i, (sid, a) in enumerate(((sid, a) for sid, d in drafts.items() for a in d.ambiguities), start=1)]
    _sheet(wb, "Requirement gaps", ["ID", "Story", "Severity", "Rule", "Requirement text", "Why it can't be tested",
                                    "Question for product", "Status"], gap_rows, [9, 8, 9, 22, 45, 45, 45, 8])
    history = []
    if previous and Path(previous).exists():
        history = [list(r) for r in load_workbook(previous)["Change log"].iter_rows(min_row=2, values_only=True)]
    history.append([version, date.today().isoformat(), prd_title,
                    f"{len(cases)} test cases drafted, {len(gap_rows)} requirement gaps raised, "
                    f"{sum(not t.cases for t in trace)} criteria without a test case, "
                    f"{sum(t.coverage == 'assumed' for t in trace)} tested on an assumption"])
    _sheet(wb, "Change log", ["Version", "Date", "Source", "Change"], history, [9, 12, 45, 70])
    wb.move_sheet("Change log", offset=-3)
    if isinstance(path, (str, Path)):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def jira_csv(stories: list[Story], gaps: list[Gap], epic: str) -> str:
    """Columns match Jira Cloud's CSV importer (map Summary, Issue Type, Description, Priority, Labels)."""
    prio = {"high": "High", "medium": "Medium", "low": "Low"}
    rows = [["Summary", "Issue Type", "Description", "Priority", "Labels"]]
    for s in stories:
        rows.append([f"{s.id} {s.title}", "Story",
                     s.narrative + "\n\nAcceptance criteria:\n" + "\n".join(f"* {c}" for c in s.criteria), "Medium", epic])
    for g in gaps:
        rows.append([f"[Spec gap] {g.story} ({g.rule.replace('_', ' ')}): {g.text[:70]}", "Task",
                     f"Rule: {g.rule}\nWhy: {g.why}\nQuestion: {g.question}", prio[g.severity], f"{epic} requirement-gap"])
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue()


def review_report(prd_title: str, stories: list[Story], drafts: dict[str, Drafted], gaps: list[Gap]) -> str:
    blocking = [g for g in gaps if g.severity == "high"]
    trace = traceability(stories, drafts, gaps)
    lines = [f"# Spec review: {prd_title}", "",
             f"**Verdict:** {'NOT ready for development' if blocking else 'Ready for development'}: "
             f"{len(blocking)} blocking (high) gap(s), {len(gaps) - len(blocking)} other.", "",
             "| Story | Criteria | Draft test cases | Criteria without a test | Gaps (high) |", "|---|---|---|---|---|"]
    for s in stories:
        sg = [g for g in gaps if g.story == s.id]
        n = len(drafts[s.id].test_cases) if s.id in drafts else 0
        uncovered = sum(not t.cases for t in trace if t.story == s.id) if s.id in drafts else "-"
        lines.append(f"| {s.id} {s.title} | {len(s.criteria)} | {n} | {uncovered} | "
                     f"{len(sg)} ({sum(g.severity == 'high' for g in sg)}) |")
    lines += ["", "## Questions for product (blocking first)", ""]
    for g in sorted(gaps, key=lambda g: {"high": 0, "medium": 1, "low": 2}[g.severity]):
        lines.append(f"- **{g.story} [{g.severity}]** \"{g.text}\": {g.why} → *{g.question}*")
    return "\n".join(lines) + "\n"


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")[:60]


def playwright_skeletons(prd_title: str, stories: list[Story], drafts: dict[str, Drafted]) -> str:
    """pytest-playwright functions, one per drafted test case, with the steps as comments to automate.

    A starting point for an automation engineer, not finished tests: locators and assertions depend on the real UI.
    """
    out = [f'"""Playwright test skeletons generated by SpecCheck from: {prd_title}.',
           "", "Fill in locators and assertions; each function's docstring is the manual test case it automates.",
           '"""', "", "import pytest", "from playwright.sync_api import Page, expect", "",
           'BASE_URL = "http://localhost:8000"', ""]
    for s in stories:
        for n, tc in enumerate(drafts.get(s.id, Drafted(test_cases=[])).test_cases, start=1):
            out += _playwright_case(case_id(s, n), s, tc)
    return "\n".join(out) + "\n"


def _playwright_case(cid: str, story: Story, tc: TestCase) -> list[str]:
    fn = f"test_{_slug(cid)}_{_slug(tc.title)}"
    body = ["", "", f"@pytest.mark.{tc.priority}", f"def {fn}(page: Page):",
            f'    """{cid} ({story.id}, covers AC {", ".join(map(str, tc.covers)) or "-"}, {tc.type}): {tc.title}',
            "", f"    Preconditions: {tc.preconditions}", f"    Expected: {tc.expected}", '    """',
            "    page.goto(BASE_URL)"]
    body += [f"    # {i}. {step}" for i, step in enumerate(tc.steps, start=1)]
    body += [f"    # Expected: {tc.expected}", "    pytest.skip(\"skeleton: add locators and assertions\")"]
    return body
