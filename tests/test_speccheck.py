import io
import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from speccheck.export import excel_library, jira_csv, playwright_skeletons, review_report, traceability
from speccheck.generate import DEFECTS, draft, review
from speccheck.lint import check, lint
from speccheck.parse import parse_prd, parse_text
from speccheck.pipeline import run

ROOT = Path(__file__).resolve().parents[1]
PRD = ROOT / "samples" / "PRD-1.1_cleaning_log_and_rejection.md"


def fake_llm(prompt: str) -> str:
    return json.dumps({"test_cases": [
        {"title": "Happy path", "type": "positive", "covers": [1], "preconditions": "QA logged in",
         "steps": ["Open record", "Act"], "expected": "Status changes; audit event", "priority": "high"},
        {"title": "Wrong role", "type": "permission", "covers": [2], "preconditions": "Operator logged in",
         "steps": ["Try the action"], "expected": "Refused with message", "priority": "medium"}],
        "ambiguities": []})


def test_parse_finds_stories_criteria_and_other_id_prefixes():
    title, stories = parse_prd(PRD)
    assert title.startswith("PRD 1.1")
    assert [s.id for s in stories] == ["US-101", "US-102", "US-103", "US-104", "US-105"]
    assert len(stories[0].criteria) == 3 and stories[0].narrative.startswith("As an operator")
    _, capa = parse_prd(ROOT / "samples" / "PRD-2.0_capa_management.md")
    assert capa[0].id == "CAPA-201" and len(capa) == 6


@pytest.mark.parametrize("text, rule", [
    ("QA can override the check as needed.", "vague_term"),
    ("QA can override the check as needed.", "gxp_control_missing"),
    ("The report layout is TBD.", "placeholder"),
    ("The system should send an email.", "weak_modal"),
    ("Sessions expire after a period of inactivity.", "missing_value"),
    ("It must be linked to the record.", "ambiguous_reference"),
    ("Results are approved before reporting.", "missing_actor"),
])
def test_each_rule_fires(text, rule):
    assert rule in {f.rule for f in check(text)}


@pytest.mark.parametrize("text", [
    "The page loads quickly (p95 under 2 seconds).",                       # vague word made measurable
    "Released records cannot be changed by any user.",                      # a prohibition is the control
    "The change request is approved by QA before implementation.",          # 'change' as a noun, actor named
    "A CAPA cannot be closed while any action is open.",                    # system rule, not a human action
    "Each action has exactly one assignee, a description and a due date.",  # 'due date' is a field
    "A correction requires a reason of at least 10 characters and an e-signature.",
])
def test_well_written_requirements_are_not_flagged(text):
    assert check(text) == []


def test_gxp_rule_uses_the_whole_story_as_context():
    _, stories = parse_text("### US-1 Reject\n- QA can reject a batch in review.\n- Rejection requires a reason and an e-signature.")
    assert not any(g.rule == "gxp_control_missing" for g in lint(stories))


def test_story_level_rules():
    _, stories = parse_text("### US-1 Thin\n- One criterion.\n### US-2 Happy only\n- Records a value.\n- Shows the value.")
    rules = {(g.story, g.rule) for g in lint(stories)}
    assert ("US-1", "thin_acceptance_criteria") in rules and ("US-2", "no_negative_behaviour") in rules


def test_eval_sets_use_known_defect_types():
    for name in ("requirements.yaml", "heldout.yaml"):
        items = yaml.safe_load((ROOT / "evals" / name).read_text(encoding="utf-8"))["items"]
        assert items and all(set(i["gold"]) <= set(DEFECTS) for i in items)


def test_llm_review_keeps_only_known_types():
    reply = json.dumps({"results": [{"i": 0, "defects": ["vague_term", "made_up"]}, {"i": 1, "defects": []}]})
    assert review(["x is fast", "y"], llm=lambda p: reply) == [["vague_term"], []]


def test_traceability_flags_criteria_without_tests():
    title, stories = parse_prd(PRD)
    drafts = {s.id: draft(s, llm=fake_llm) for s in stories}
    trace = traceability(stories, drafts, lint(stories))
    us101 = [t for t in trace if t.story == "US-101"]
    assert us101[0].cases == ["TC-101-01"] and us101[1].cases == ["TC-101-02"] and us101[2].cases == []


def test_excel_library_versions_and_sheets(tmp_path):
    title, stories = parse_prd(PRD)
    drafts = {s.id: draft(s, llm=fake_llm) for s in stories}
    gaps = lint(stories)
    v1 = excel_library(tmp_path / "test_library_v1.0.xlsx", "1.0", title, stories, drafts, gaps)
    v2 = excel_library(tmp_path / "test_library_v1.1.xlsx", "1.1", title, stories, drafts, gaps, previous=v1)
    wb = load_workbook(v2)
    assert wb.sheetnames == ["Change log", "Test cases", "Traceability", "Requirement gaps"]
    assert [r[0] for r in wb["Change log"].iter_rows(min_row=2, values_only=True)] == ["1.0", "1.1"]
    buf = io.BytesIO()
    excel_library(buf, "1.0", title, stories, drafts, gaps)
    assert buf.getvalue()[:2] == b"PK"  # the web app writes to memory


def test_jira_report_and_playwright_outputs_are_valid():
    title, stories = parse_prd(PRD)
    drafts = {s.id: draft(s, llm=fake_llm) for s in stories}
    gaps = lint(stories)
    csv_text = jira_csv(stories, gaps, "PRD-1.1")
    assert csv_text.splitlines()[0] == "Summary,Issue Type,Description,Priority,Labels"
    assert csv_text.count(",Story,") == 5 and "requirement-gap" in csv_text
    assert "NOT ready for development" in review_report(title, stories, drafts, gaps)
    code = playwright_skeletons(title, stories, drafts)
    compile(code, "skeletons.py", "exec")
    assert code.count("def test_") == 10


def test_pipeline_survives_a_failing_llm():
    def broken(prompt: str) -> str:
        raise RuntimeError("503")
    r = run(PRD.read_text(encoding="utf-8"), use_llm=True, llm=broken)
    assert r.gaps and not r.drafts and set(r.skipped) == {"US-101", "US-102", "US-103", "US-104", "US-105"}
